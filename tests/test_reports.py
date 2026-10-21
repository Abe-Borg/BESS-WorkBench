import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from bess_workbench.controller import ProjectController
from bess_workbench.reports import export_all
from bess_workbench.ruleset import load_rules
from bess_workbench.sample_data import build_sample_project


def test_sample_exports_are_valid(tmp_path):
    rules = load_rules()
    controller = ProjectController(build_sample_project(rules), rules)
    controller.project.ai_narrative = "The fire area is at the cap."
    controller.project.ai_narrative_model = "claude-sonnet-5-5"
    paths = export_all(controller.project, controller.analysis, tmp_path)
    docx_xml = zipfile.ZipFile(paths["docx"]).read("word/document.xml").decode("utf-8")
    assert "Sample - 4 Room UPS BESS" in docx_xml
    assert "Design aid. Verify values against adopted codes and AHJ." in docx_xml
    assert "AI-drafted" in docx_xml
    assert "Arial" in docx_xml
    book = load_workbook(paths["xlsx"])
    assert {"Cover", "Rooms", "ESS", "Checks", "Points", "Cause and Effect", "RFIs"} <= set(book.sheetnames)
    assert book["Cover"]["A3"].value.startswith("Design aid")
    reader = PdfReader(paths["pdf"])
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Sample - 4 Room UPS BESS" in text
    assert "Design aid" in text
    assert paths["cause_effect"].stat().st_size > 1000
    assert Path(paths["pdf"]).read_bytes().startswith(b"%PDF")
