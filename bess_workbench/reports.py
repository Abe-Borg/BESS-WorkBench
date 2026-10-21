"""Basis of Design DOCX, workbook, summary PDF, and cause-and-effect export."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from bess_workbench import __version__
from bess_workbench.analysis import Analysis
from bess_workbench.calcs import CE_COLUMNS
from bess_workbench.formatting import fmt_fixed, fmt_num
from bess_workbench.models import Project, display_rfi_id

DISCLAIMER = "Design aid. Verify values against adopted codes and AHJ."
NAVY = "1F4E79"
STATUS_COLORS = {
    "PASS": "1E7F4F",
    "FAIL": "B42318",
    "WARN": "B54708",
    "BLOCKED": "7A271A",
    "INFO": "175CD3",
}


def find_body_font() -> tuple[str, str | None, str | None, bool]:
    """Return reportlab font name, regular path, bold path, and whether Arial was found."""
    pairs = [
        ("Arial", "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", True),
        ("Arial", "/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf", True),
        ("Arial", "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf", "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf", True),
        (
            "LiberationSans",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
            False,
        ),
        (
            "DejaVuSans",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            False,
        ),
    ]
    for name, regular, bold, is_arial in pairs:
        if Path(regular).is_file():
            return name, regular, bold if Path(bold).is_file() else None, is_arial
    return "Helvetica", None, None, False


def _register_pdf_fonts() -> tuple[str, str, bool]:
    name, regular, bold, is_arial = find_body_font()
    if regular and name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, regular))
        if bold:
            pdfmetrics.registerFont(TTFont(name + "-Bold", bold))
        else:
            pdfmetrics.registerFont(TTFont(name + "-Bold", regular))
    bold_name = "Helvetica-Bold" if name == "Helvetica" else name + "-Bold"
    return name, bold_name, is_arial


def _set_run_font(run, size: float = 11, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Arial"
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Arial")


def _style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), "Arial")
    for style_name in ("Heading 1", "Heading 2", "Title"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        _set_run_font(run, 10, bold=True, color=RGBColor(255, 255, 255))
        _shade(cell, NAVY)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run("" if value is None else str(value))
            _set_run_font(run, 9)
    document.add_paragraph()


def _para(document: Document, text: str, *, bold: bool = False, size: float = 11) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _set_run_font(run, size, bold=bold)


def write_basis_of_design(project: Project, analysis: Analysis, path: str | Path) -> Path:
    target = Path(path)
    document = Document()
    _style_document(document)
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run(f"BESS WorkBench {__version__}  |  {DISCLAIMER}")
    _set_run_font(run, 9, color=RGBColor(0x7A, 0x4E, 0x00))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(DISCLAIMER)
    _set_run_font(footer_run, 8)

    title = document.add_heading("Basis of Design", level=0)
    for run in title.runs:
        _set_run_font(run, 20, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    _para(document, project.meta.name, bold=True, size=14)
    if project.meta.synthetic:
        _para(document, "SYNTHETIC PROJECT — not a real facility.", bold=True)
    _para(document, DISCLAIMER)

    document.add_heading("Project", level=1)
    _add_table(
        document,
        ["Field", "Value"],
        [
            ["Project number", project.meta.project_number],
            ["Client", project.meta.client],
            ["Site", project.meta.site],
            ["Prepared by", project.meta.prepared_by],
            ["Ruleset", f"{project.ruleset_id} {project.ruleset_version}"],
            ["Occupancy", project.code_basis.occupancy],
            ["Jurisdiction", project.code_basis.jurisdiction],
        ],
    )
    if project.meta.notes:
        _para(document, project.meta.notes)
    document.add_heading("Code basis", level=1)
    basis = project.code_basis
    _add_table(
        document,
        ["Document", "Edition"],
        [
            ["IFC", basis.ifc],
            ["IBC", basis.ibc],
            ["IMC", basis.imc],
            ["NEC", basis.nec],
            ["NFPA 855", basis.nfpa_855],
            ["NFPA 72", basis.nfpa_72],
            ["NFPA 13", basis.nfpa_13],
            ["NFPA 69", basis.nfpa_69],
        ],
    )
    _para(document, "AHJ notes: " + (basis.ahj_notes or "None recorded."))

    document.add_heading("Narrative", level=1)
    if project.ai_narrative.strip():
        label = "AI-drafted and edited by the user" if project.ai_narrative_edited else "AI-drafted"
        _para(
            document,
            f"{label}. This narrative is not a calculation and does not change any pass/fail result. "
            f"Model: {project.ai_narrative_model or 'unspecified'}.",
            bold=True,
        )
        _para(document, project.ai_narrative)
    else:
        _para(document, "No AI narrative has been drafted.")

    document.add_heading("Rooms", level=1)
    _add_table(
        document,
        ["Room", "Area sf", "Height ft", "Wall h", "Floor h", "Ceiling h", "Fire area", "Smoke"],
        [
            [
                room.name,
                fmt_num(room.area_sf),
                fmt_num(room.height_ft),
                fmt_num(room.wall_rating_hr),
                fmt_num(room.floor_rating_hr),
                fmt_num(room.ceiling_rating_hr),
                room.fire_area_id,
                "yes" if room.smoke_detection else "no",
            ]
            for room in project.rooms
        ],
    )
    document.add_heading("ESS inventory", level=1)
    _add_table(
        document,
        ["Group", "Room", "Chemistry", "kWh", "Spacing in", "UL 9540", "UL 9540A"],
        [
            [
                group.name,
                group.room_id,
                group.chemistry,
                fmt_num(group.kwh),
                "" if group.spacing_in is None else fmt_num(group.spacing_in),
                "yes" if group.ul_9540_listed else "no",
                "yes" if group.ul_9540a_report else "no",
            ]
            for group in project.ess_groups
        ],
    )
    document.add_heading("Fire areas", level=1)
    _add_table(
        document,
        ["Fire area", "kWh", "Barrier h", "Floor/ceiling h", "Separation", "Energy"],
        [
            [
                item.fire_area_id,
                fmt_num(item.li_ion_kwh),
                fmt_num(item.barrier_hr),
                fmt_num(item.horizontal_hr),
                "yes" if item.separation_valid else "no",
                item.energy_status.value if item.energy_checked else "combined elsewhere",
            ]
            for item in analysis.rollups
        ],
    )
    document.add_heading("Sprinklers", level=1)
    summary = analysis.sprinkler_summary
    _para(document, summary.formula)
    _para(
        document,
        f"System {summary.system_type}. Governing {summary.governing_room_id or 'n/a'}: "
        f"{fmt_fixed(summary.governing_sprinkler_gpm)} gpm sprinkler + {fmt_fixed(summary.hose_gpm)} gpm hose "
        f"= {fmt_fixed(summary.total_gpm)} gpm for {fmt_num(summary.duration_min)} min "
        f"({fmt_num(summary.volume_gal)} gal worksheet). C={fmt_num(summary.c_factor)}.",
    )
    _para(document, "Duct sprinkler: " + analysis.duct.formula)
    document.add_heading("Gas detection", level=1)
    _para(
        document,
        f"Monitoring-only: {'yes' if project.gas.monitoring_only else 'no'}. "
        f"Add alternate: {'yes' if project.gas.add_alternate else 'no'}. "
        f"Alarm {fmt_num(project.gas.alarm_pct_lfl)}% LFL. Design {fmt_num(project.gas.design_pct_lfl)}% LFL. "
        f"Proof points: {analysis.proof_required}, on the gas panel {analysis.proof_on_panel}, on FACP modules {analysis.proof_on_facp}.",
    )
    document.add_heading("Batteries", level=1)
    _para(document, "FACP: " + analysis.battery_facp.formula)
    _para(document, "Gas panel: " + analysis.battery_gas.formula)
    document.add_heading("Open items", level=1)
    open_items = [item for item in project.rfis if item.status == "open"]
    _add_table(
        document,
        ["ID", "Status", "Assignee", "Title", "Text"],
        [
            [display_rfi_id(item), item.status, item.assignee, item.title, item.text]
            for item in open_items
        ]
        or [["—", "—", "—", "None open", ""]],
    )
    if project.provenance:
        document.add_heading("Accepted document fields", level=1)
        _add_table(
            document,
            ["Field", "Document", "Page", "Detail"],
            [[item.path, item.source_document, "" if item.page is None else str(item.page), item.detail] for item in project.provenance],
        )
    document.add_heading("Checks", level=1)
    _add_table(
        document,
        ["Status", "ID", "Value", "Citation", "Tags"],
        [[item.status.value, item.id, item.value, item.citation, ", ".join(item.tags)] for item in analysis.checks],
    )
    document.save(target)
    return target


def _xlsx_header(ws, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(1, index, header)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 4, 14), 42)


def _xlsx_body(ws, rows: list[list]) -> None:
    font = Font(name="Arial", size=10)
    align = Alignment(wrap_text=True, vertical="top")
    for row_index, values in enumerate(rows, start=2):
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col_index, value)
            cell.font = font
            cell.alignment = align
            if col_index == 1 and isinstance(value, str) and value in STATUS_COLORS:
                cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[value])
                cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)


def write_workbook(project: Project, analysis: Analysis, path: str | Path) -> Path:
    target = Path(path)
    book = Workbook()
    cover = book.active
    cover.title = "Cover"
    cover["A1"] = "BESS WorkBench"
    cover["A1"].font = Font(name="Arial", size=18, bold=True, color=NAVY)
    cover["A2"] = project.meta.name
    cover["A2"].font = Font(name="Arial", size=14, bold=True)
    cover["A3"] = DISCLAIMER
    cover["A3"].font = Font(name="Arial", size=11, color="7A4E00")
    cover["A4"] = "SYNTHETIC PROJECT" if project.meta.synthetic else ""
    cover["A5"] = project.meta.notes
    cover["A5"].alignment = Alignment(wrap_text=True)
    cover.column_dimensions["A"].width = 110
    cover.row_dimensions[5].height = 60

    rooms = book.create_sheet("Rooms")
    room_headers = ["ID", "Name", "Area sf", "Height ft", "Wall h", "Floor h", "Ceiling h", "Fire area", "Smoke"]
    _xlsx_header(rooms, room_headers)
    _xlsx_body(
        rooms,
        [
            [r.id, r.name, r.area_sf, r.height_ft, r.wall_rating_hr, r.floor_rating_hr, r.ceiling_rating_hr, r.fire_area_id, r.smoke_detection]
            for r in project.rooms
        ],
    )
    ess = book.create_sheet("ESS")
    _xlsx_header(ess, ["ID", "Name", "Room", "Chemistry", "kWh", "Spacing in", "UL 9540", "UL 9540A"])
    _xlsx_body(
        ess,
        [
            [g.id, g.name, g.room_id, g.chemistry, g.kwh, g.spacing_in, g.ul_9540_listed, g.ul_9540a_report]
            for g in project.ess_groups
        ],
    )
    areas = book.create_sheet("Fire Areas")
    _xlsx_header(areas, ["ID", "Name", "Rooms", "Li-ion kWh", "Barrier h", "Horizontal h", "Separation valid", "Energy status", "Note"])
    _xlsx_body(
        areas,
        [
            [
                item.fire_area_id,
                item.name,
                ", ".join(item.room_ids),
                item.li_ion_kwh,
                item.barrier_hr,
                item.horizontal_hr,
                item.separation_valid,
                item.energy_status.value if item.energy_checked else "",
                item.note,
            ]
            for item in analysis.rollups
        ],
    )
    checks = book.create_sheet("Checks")
    _xlsx_header(checks, ["Status", "ID", "Screen", "Title", "Value", "Message", "Citation", "Edition", "Tags"])
    _xlsx_body(
        checks,
        [
            [c.status.value, c.id, c.screen, c.title, c.value, c.message, c.citation, c.edition, ", ".join(c.tags)]
            for c in analysis.checks
        ],
    )
    points = book.create_sheet("Points")
    _xlsx_header(points, ["Point", "Description", "Kind", "Landing", "Normal", "Off-normal"])
    _xlsx_body(points, [[p.point_id, p.description, p.kind, p.landing, p.normal_state, p.offnormal_state] for p in analysis.points])
    ce = book.create_sheet("Cause and Effect")
    _write_ce_sheet(ce, project)
    sprinklers = book.create_sheet("Sprinklers")
    _xlsx_header(sprinklers, ["Room", "Area sf", "Base sf", "Increase sf", "Design sf", "Density", "gpm", "Governing"])
    _xlsx_body(
        sprinklers,
        [
            [row.room_name, row.room_area_sf, row.base_area_sf, row.increase_sf, row.design_area_sf, row.density_gpm_sf, row.sprinkler_gpm, row.governing]
            for row in analysis.sprinkler_rows
        ],
    )
    battery = book.create_sheet("Battery")
    _xlsx_header(battery, ["Name", "Raw Ah", "Margin", "Aging", "Temp", "Required Ah", "Next standard Ah", "Formula"])
    _xlsx_body(
        battery,
        [
            [item.name, item.raw_ah, item.margin_used, item.aging_used, item.temp_used, item.required_ah, item.next_standard_ah, item.formula]
            for item in (analysis.battery_facp, analysis.battery_gas)
        ],
    )
    rfis = book.create_sheet("RFIs")
    _xlsx_header(rfis, ["ID", "Status", "Source", "Assignee", "Due", "Rule", "Title", "Text", "AI drafted"])
    _xlsx_body(
        rfis,
        [
            [display_rfi_id(item), item.status, item.source, item.assignee, item.due, item.rule_id, item.title, item.text, item.ai_drafted]
            for item in project.rfis
        ],
    )
    book.save(target)
    return target


def _write_ce_sheet(ws, project: Project) -> None:
    headers = ["Cause", *CE_COLUMNS, "Notes"]
    _xlsx_header(ws, headers)
    rows = []
    for row in project.cause_effect:
        rows.append([row.cause, *[row.effects.get(column, "") for column in CE_COLUMNS], row.notes])
    _xlsx_body(ws, rows)


def write_cause_effect(project: Project, path: str | Path) -> Path:
    target = Path(path)
    book = Workbook()
    sheet = book.active
    sheet.title = "Cause and Effect"
    _write_ce_sheet(sheet, project)
    book.save(target)
    return target


def write_summary_pdf(project: Project, analysis: Analysis, path: str | Path) -> Path:
    target = Path(path)
    font, bold, is_arial = _register_pdf_fonts()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyA", fontName=font, fontSize=9, leading=12, textColor=colors.HexColor("#1C2833")))
    styles.add(ParagraphStyle(name="H1A", fontName=bold, fontSize=14, leading=18, textColor=colors.HexColor("#1F4E79"), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="SmallA", fontName=font, fontSize=8, leading=10, textColor=colors.HexColor("#7A4E00")))
    story = []
    story.append(Paragraph("BESS WorkBench — Summary", styles["H1A"]))
    story.append(Paragraph(DISCLAIMER, styles["SmallA"]))
    if not is_arial:
        story.append(Paragraph(f"PDF body font is {font} because Arial is not installed on this machine. The DOCX requests Arial.", styles["SmallA"]))
    story.append(Spacer(1, 6))
    story.append(Paragraph(project.meta.name, styles["H1A"]))
    if project.meta.synthetic:
        story.append(Paragraph("SYNTHETIC PROJECT — not a real facility.", styles["BodyA"]))
    counts: dict[str, int] = {}
    for check in analysis.checks:
        counts[check.status.value] = counts.get(check.status.value, 0) + 1
    count_text = " · ".join(f"{counts.get(name, 0)} {name}" for name in ("BLOCKED", "FAIL", "WARN", "INFO", "PASS"))
    story.append(Paragraph(count_text, styles["BodyA"]))
    story.append(Paragraph(f"Ruleset {project.ruleset_id} {project.ruleset_version}", styles["BodyA"]))
    story.append(Paragraph("Fire areas", styles["H1A"]))
    data = [["Fire area", "Li-ion kWh", "Status", "Note"]]
    for item in analysis.rollups:
        data.append(
            [
                item.fire_area_id,
                fmt_num(item.li_ion_kwh),
                item.energy_status.value if item.energy_checked else "—",
                item.note,
            ]
        )
    story.append(_pdf_table(data, font, bold))
    story.append(Paragraph("Sprinklers and batteries", styles["H1A"]))
    story.append(Paragraph(analysis.sprinkler_summary.formula, styles["BodyA"]))
    story.append(Paragraph(analysis.duct.formula, styles["BodyA"]))
    story.append(Paragraph("FACP: " + analysis.battery_facp.formula, styles["BodyA"]))
    story.append(Paragraph("Gas panel: " + analysis.battery_gas.formula, styles["BodyA"]))
    story.append(Paragraph("Open items", styles["H1A"]))
    open_rows = [["ID", "Assignee", "Title"]]
    for item in project.rfis:
        if item.status != "open":
            continue
        open_rows.append([display_rfi_id(item), item.assignee, item.title])
    if len(open_rows) == 1:
        open_rows.append(["—", "—", "None"])
    story.append(_pdf_table(open_rows, font, bold))
    story.append(Spacer(1, 8))
    story.append(Paragraph(DISCLAIMER, styles["SmallA"]))

    def _header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#1F4E79"))
        canvas.rect(0, letter[1] - 28, letter[0], 28, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(font, 8)
        canvas.drawString(36, letter[1] - 18, f"BESS WorkBench {__version__}")
        canvas.drawRightString(letter[0] - 36, letter[1] - 18, DISCLAIMER)
        canvas.setFillColor(colors.HexColor("#5D6D7E"))
        canvas.drawString(36, 22, project.meta.name[:80])
        canvas.drawRightString(letter[0] - 36, 22, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(target), pagesize=letter, topMargin=42, bottomMargin=40, leftMargin=36, rightMargin=36, title=project.meta.name)
    doc.build(story, onFirstPage=_header_footer, onLaterPages=_header_footer)
    return target


def _pdf_table(data: list[list[str]], font: str, bold: str) -> Table:
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), bold),
                ("FONTNAME", (0, 1), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7F9FB")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def export_all(project: Project, analysis: Analysis, outdir: str | Path) -> dict[str, Path]:
    folder = Path(outdir)
    folder.mkdir(parents=True, exist_ok=True)
    paths = {
        "docx": write_basis_of_design(project, analysis, folder / "Basis-of-Design.docx"),
        "xlsx": write_workbook(project, analysis, folder / "BESS-Workbook.xlsx"),
        "pdf": write_summary_pdf(project, analysis, folder / "Summary.pdf"),
        "cause_effect": write_cause_effect(project, folder / "Cause-and-Effect.xlsx"),
    }
    return paths
